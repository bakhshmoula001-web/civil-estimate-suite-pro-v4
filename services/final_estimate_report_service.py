"""
Civil Estimate Suite Pro v4.0
Final Estimate Excel / PDF Report Service
"""

from __future__ import annotations

from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
)


class FinalEstimateReportService:
    """Generate consolidated project estimate reports."""

    def __init__(self, export_folder="exports"):
        self.export_folder = Path(export_folder)
        self.export_folder.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _num(value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _money(value):
        return f"PKR {float(value or 0):,.2f}"

    @staticmethod
    def _project_value(project, field):
        return str(getattr(project, field, "") or "").strip()

    @classmethod
    def _dimensions(cls, item):
        """Return the saved engineering dimensions without losing calculator detail."""
        direct = (
            getattr(item, "dimensions", None)
            or getattr(item, "dimension", None)
            or getattr(item, "dimensions_text", None)
        )
        if direct:
            return str(direct).strip()

        parts = []
        fields = [
            ("L", ("length", "item_length")),
            ("W", ("width", "item_width")),
            ("H", ("height", "item_height")),
            ("T", ("thickness", "item_thickness", "depth")),
            ("D", ("depth", "item_depth")),
        ]
        used = set()
        for label, names in fields:
            for name in names:
                value = getattr(item, name, None)
                if value not in (None, ""):
                    text = str(value).strip()
                    if text and name not in used:
                        parts.append(f"{label}={text}")
                        used.add(name)
                        break

        unit = getattr(item, "dimension_unit", None) or getattr(item, "unit", "")
        if parts:
            return " × ".join(parts) + (f" {unit}" if unit else "")
        return "—"

    @classmethod
    def _boq_rows(cls, items):
        rows = []
        for serial, item in enumerate(items or [], start=1):
            rows.append({
                "sno": serial,
                "item_no": str(getattr(item, "item_no", "") or ""),
                "description": str(getattr(item, "description", "") or ""),
                "dimensions": cls._dimensions(item),
                "unit": str(getattr(item, "unit", "") or ""),
                "quantity": cls._num(getattr(item, "quantity", 0)),
                "rate": cls._num(getattr(item, "rate", 0)),
                "amount": cls._num(getattr(item, "amount", 0)),
            })
            if rows[-1]["amount"] == 0:
                rows[-1]["amount"] = round(
                    rows[-1]["quantity"] * rows[-1]["rate"], 2
                )
        return rows

    def _summary(self, items, analyses):
        material = skilled = unskilled = other = total = 0.0
        materials = {}
        labour = []

        for item in items or []:
            analysis = (analyses or {}).get(getattr(item, "id", None))
            if not isinstance(analysis, dict):
                continue

            material += self._num(
                analysis.get("material_cost", analysis.get("material_total", 0))
            )
            skilled += self._num(analysis.get("skilled_labour_cost", 0))
            unskilled += self._num(analysis.get("unskilled_labour_cost", 0))
            other += self._num(analysis.get("other_labour_cost", 0))

            labour_total = self._num(
                analysis.get(
                    "labour_cost",
                    analysis.get("labour_total", skilled + unskilled + other),
                )
            )
            total += self._num(
                analysis.get(
                    "total_cost",
                    self._num(
                        analysis.get(
                            "material_cost",
                            analysis.get("material_total", 0),
                        )
                    ) + labour_total,
                )
            )

            for row in analysis.get("materials", []) or []:
                name = str(
                    row.get("name")
                    or row.get("material_name")
                    or "Material"
                ).strip()
                qty = self._num(row.get("quantity", row.get("QTY")))
                unit = str(row.get("unit", row.get("UNIT", "")) or "")
                amount = self._num(
                    row.get("amount", row.get("cost", qty * self._num(row.get("rate"))))
                )
                key = (name, unit)
                if key not in materials:
                    materials[key] = {"quantity": 0.0, "amount": 0.0}
                materials[key]["quantity"] += qty
                materials[key]["amount"] += amount

            for row in analysis.get("labour", []) or []:
                name = str(
                    row.get("name")
                    or row.get("labour")
                    or "Labour"
                ).strip()
                unit = str(row.get("unit", "day") or "day")
                qty = self._num(row.get("quantity", row.get("days")))
                amount = self._num(
                    row.get("amount", qty * self._num(row.get("rate")))
                )
                labour.append((name, unit, qty, amount))

        labour_total = skilled + unskilled + other
        if total == 0:
            total = material + labour_total

        return {
            "material": round(material, 2),
            "skilled": round(skilled, 2),
            "unskilled": round(unskilled, 2),
            "other": round(other, 2),
            "labour": round(labour_total, 2),
            "total": round(total, 2),
            "materials": materials,
            "labour_rows": labour,
        }

    def export_excel(self, project, items, analyses, file_path=None):
        if project is None:
            raise ValueError("No active project selected.")

        path = Path(file_path) if file_path else (
            self.export_folder
            / f"{self._project_value(project, 'project_code') or 'PROJECT'}_Final_Estimate.xlsx"
        )
        path.parent.mkdir(parents=True, exist_ok=True)

        summary = self._summary(items, analyses)

        wb = Workbook()
        ws = wb.active
        ws.title = "Final Estimate"
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:H1")
        ws["A1"] = "FINAL PROJECT ESTIMATE"
        ws["A1"].font = Font(size=18, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")

        project_fields = [
            ("Project Code", self._project_value(project, "project_code")),
            ("Project Name", self._project_value(project, "project_name")),
            ("Client", self._project_value(project, "client_name")),
            ("Location", self._project_value(project, "location")),
        ]

        row = 3
        for label, value in project_fields:
            ws.cell(row, 1, label).font = Font(bold=True)
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
            ws.cell(row, 2, value)
            row += 1

        row += 1
        ws.cell(row, 1, "COST SUMMARY").font = Font(size=13, bold=True)
        row += 1

        summary_rows = [
            ("Material Cost", summary["material"]),
            ("Skilled Labour Cost", summary["skilled"]),
            ("Unskilled Labour Cost", summary["unskilled"]),
            ("Other Labour Cost", summary["other"]),
            ("Total Labour Cost", summary["labour"]),
            ("GRAND TOTAL", summary["total"]),
        ]

        for label, value in summary_rows:
            ws.cell(row, 1, label).font = Font(bold=True)
            ws.cell(row, 7, value)
            ws.cell(row, 7).number_format = '#,##0.00'
            row += 1

        row += 1
        ws.cell(row, 1, "MATERIAL REQUIREMENT").font = Font(size=13, bold=True)
        row += 1
        headers = ["Material", "Unit", "Quantity", "Amount (PKR)"]
        for col, value in enumerate(headers, 1):
            c = ws.cell(row, col, value)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")

        row += 1
        for (name, unit), data in sorted(summary["materials"].items()):
            ws.cell(row, 1, name)
            ws.cell(row, 2, unit)
            ws.cell(row, 3, round(data["quantity"], 3))
            ws.cell(row, 4, round(data["amount"], 2))
            ws.cell(row, 3).number_format = '#,##0.000'
            ws.cell(row, 4).number_format = '#,##0.00'
            row += 1

        row += 1
        ws.cell(row, 1, "LABOUR REQUIREMENT").font = Font(size=13, bold=True)
        row += 1
        for col, value in enumerate(["Labour", "Unit", "Quantity", "Amount (PKR)"], 1):
            c = ws.cell(row, col, value)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        row += 1
        for name, unit, qty, amount in summary["labour_rows"]:
            ws.cell(row, 1, name)
            ws.cell(row, 2, unit)
            ws.cell(row, 3, round(qty, 3))
            ws.cell(row, 4, round(amount, 2))
            ws.cell(row, 3).number_format = '#,##0.000'
            ws.cell(row, 4).number_format = '#,##0.00'
            row += 1

        row += 2
        ws.cell(row, 1, "Generated")
        ws.cell(row, 2, datetime.now().strftime("%d-%m-%Y %H:%M"))
        ws.cell(row, 1).font = Font(bold=True)

        widths = {1: 30, 2: 18, 3: 16, 4: 20, 5: 4, 6: 4, 7: 22, 8: 18}
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width

        thin = Side(style="thin", color="D0D0D0")
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                cell.border = Border(
                    left=thin, right=thin, top=thin, bottom=thin
                )
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal=(
                        "right" if cell.column in (3, 4, 7)
                        else "left"
                    ),
                )

        # =================================================
        # BOQ DETAIL SHEET
        # =================================================
        boq_ws = wb.create_sheet("BOQ Detail")
        boq_ws.sheet_view.showGridLines = False
        boq_ws.merge_cells("A1:H1")
        boq_ws["A1"] = "BOQ DETAIL - ENGINEERING QUANTITIES & COST"
        boq_ws["A1"].font = Font(size=16, bold=True)
        boq_ws["A1"].alignment = Alignment(horizontal="center")

        boq_headers = [
            "S.No.", "Item No.", "Description", "Dimensions",
            "Unit", "Quantity", "Rate (PKR)", "Amount (PKR)"
        ]
        header_row = 3
        for col, value in enumerate(boq_headers, 1):
            cell = boq_ws.cell(header_row, col, value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        boq_total = 0.0
        row = header_row + 1
        for data in self._boq_rows(items):
            values = [
                data["sno"], data["item_no"], data["description"],
                data["dimensions"], data["unit"], data["quantity"],
                data["rate"], data["amount"]
            ]
            for col, value in enumerate(values, 1):
                cell = boq_ws.cell(row, col, value)
                cell.alignment = Alignment(
                    horizontal="right" if col in (6, 7, 8) else (
                        "center" if col in (1, 5) else "left"
                    ),
                    vertical="top",
                    wrap_text=True,
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

            boq_ws.cell(row, 6).number_format = '#,##0.000'
            boq_ws.cell(row, 7).number_format = '#,##0.00'
            boq_ws.cell(row, 8).number_format = '#,##0.00'
            boq_total += data["amount"]
            row += 1

        boq_ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        boq_ws.cell(row, 1, "GRAND TOTAL").font = Font(bold=True, size=12)
        boq_ws.cell(row, 1).alignment = Alignment(horizontal="right")
        boq_ws.cell(row, 8, round(boq_total, 2)).font = Font(bold=True, size=12)
        boq_ws.cell(row, 8).number_format = '#,##0.00'

        widths = {
            1: 8, 2: 15, 3: 34, 4: 34, 5: 10, 6: 16, 7: 18, 8: 20
        }
        for col, width in widths.items():
            boq_ws.column_dimensions[get_column_letter(col)].width = width
        boq_ws.freeze_panes = "A4"
        boq_ws.auto_filter.ref = f"A3:H{max(row-1, 3)}"

        wb.save(path)
        return path

    def export_pdf(self, project, items, analyses, file_path=None):
        if project is None:
            raise ValueError("No active project selected.")

        path = Path(file_path) if file_path else (
            self.export_folder
            / f"{self._project_value(project, 'project_code') or 'PROJECT'}_Final_Estimate.pdf"
        )
        path.parent.mkdir(parents=True, exist_ok=True)

        summary = self._summary(items, analyses)

        doc = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
        )

        styles = getSampleStyleSheet()
        title = ParagraphStyle(
            "FinalTitle", parent=styles["Title"], alignment=TA_CENTER,
            fontSize=18, leading=22, spaceAfter=8
        )
        heading = ParagraphStyle(
            "Section", parent=styles["Heading2"], fontSize=12,
            leading=15, spaceBefore=7, spaceAfter=5
        )
        body = ParagraphStyle(
            "BodySmall", parent=styles["BodyText"], fontSize=8.5,
            leading=11
        )

        story = [
            Paragraph("FINAL PROJECT ESTIMATE", title),
            Paragraph(
                f"<b>Project:</b> {self._project_value(project, 'project_name')} &nbsp;&nbsp; "
                f"<b>Code:</b> {self._project_value(project, 'project_code')} &nbsp;&nbsp; "
                f"<b>Client:</b> {self._project_value(project, 'client_name')} &nbsp;&nbsp; "
                f"<b>Location:</b> {self._project_value(project, 'location')}",
                body,
            ),
            Spacer(1, 5 * mm),
            Paragraph("Cost Summary", heading),
        ]

        data = [
            ["Component", "Amount (PKR)"],
            ["Material Cost", f"{summary['material']:,.2f}"],
            ["Skilled Labour Cost", f"{summary['skilled']:,.2f}"],
            ["Unskilled Labour Cost", f"{summary['unskilled']:,.2f}"],
            ["Other Labour Cost", f"{summary['other']:,.2f}"],
            ["Total Labour Cost", f"{summary['labour']:,.2f}"],
            ["GRAND TOTAL", f"{summary['total']:,.2f}"],
        ]
        table = Table(data, colWidths=[85 * mm, 55 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E7E6E6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(table)

        # =================================================
        # BOQ DETAIL
        # =================================================
        story.append(Paragraph("BOQ Detail - Engineering Quantities", heading))
        bdata = [[
            "S.No.", "Item No.", "Description", "Dimensions",
            "Unit", "Quantity", "Rate (PKR)", "Amount (PKR)"
        ]]
        for data in self._boq_rows(items):
            bdata.append([
                str(data["sno"]),
                data["item_no"],
                data["description"],
                data["dimensions"],
                data["unit"],
                f"{data['quantity']:,.3f}",
                f"{data['rate']:,.2f}",
                f"{data['amount']:,.2f}",
            ])
        if len(bdata) == 1:
            bdata.append(["", "", "No BOQ items available", "", "", "", "", ""])

        bt = Table(
            bdata,
            colWidths=[12*mm, 25*mm, 48*mm, 48*mm, 16*mm, 25*mm, 27*mm, 30*mm],
            repeatRows=1,
        )
        bt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (1, -1), "CENTER"),
            ("ALIGN", (4, 1), (4, -1), "CENTER"),
            ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ]))
        story.append(bt)

        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Material Requirement", heading))
        mdata = [["Material", "Unit", "Quantity", "Amount (PKR)"]]
        for (name, unit), d in sorted(summary["materials"].items()):
            mdata.append([
                name, unit, f"{d['quantity']:,.3f}", f"{d['amount']:,.2f}"
            ])
        if len(mdata) == 1:
            mdata.append(["No material schedule available", "", "", ""])
        mt = Table(mdata, colWidths=[70*mm, 25*mm, 35*mm, 45*mm], repeatRows=1)
        mt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(mt)

        story.append(Paragraph("Labour Requirement", heading))
        ldata = [["Labour", "Unit", "Quantity", "Amount (PKR)"]]
        for name, unit, qty, amount in summary["labour_rows"]:
            ldata.append([name, unit, f"{qty:,.3f}", f"{amount:,.2f}"])
        if len(ldata) == 1:
            ldata.append(["No labour schedule available", "", "", ""])
        lt = Table(ldata, colWidths=[70*mm, 25*mm, 35*mm, 45*mm], repeatRows=1)
        lt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (2, 1), (-1, -1), "RIGHT"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(lt)
        story.append(Spacer(1, 5 * mm))
        story.append(
            Paragraph(
                f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}",
                body,
            )
        )

        doc.build(story)
        return path


__all__ = ["FinalEstimateReportService"]
