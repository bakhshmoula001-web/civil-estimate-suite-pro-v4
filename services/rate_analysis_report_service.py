"""
Civil Estimate Suite Pro v4.0
Rate Analysis / Unit Rate Build-up Report Service
"""

from __future__ import annotations

from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


class RateAnalysisReportService:
    """Generate item-wise material + labour + unit-rate build-up reports."""

    def __init__(self, export_folder="exports"):
        self.export_folder = Path(export_folder)
        self.export_folder.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _num(value):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0.0

    @classmethod
    def _value(cls, data, *keys):
        for key in keys:
            if key in data and data[key] not in (None, ""):
                return cls._num(data[key])
        return 0.0

    @classmethod
    def _rows(cls, items, analyses):
        rows = []
        for item in items or []:
            item_id = getattr(item, "id", None)
            analysis = (analyses or {}).get(item_id, {})
            if not isinstance(analysis, dict):
                analysis = {}

            qty = cls._num(getattr(item, "quantity", 0))
            material = cls._value(
                analysis, "material_cost", "material_total"
            )
            skilled = cls._value(
                analysis, "skilled_labour_cost", "skilled_cost"
            )
            unskilled = cls._value(
                analysis, "unskilled_labour_cost", "unskilled_cost"
            )
            other = cls._value(
                analysis, "other_labour_cost", "other_labour"
            )
            labour = cls._value(
                analysis, "labour_cost", "labour_total"
            )
            if labour == 0:
                labour = skilled + unskilled + other

            total = cls._value(analysis, "total_cost")
            if total == 0:
                total = material + labour

            unit_rate = total / qty if qty else 0.0

            rows.append({
                "item_no": str(getattr(item, "item_no", "") or ""),
                "description": str(getattr(item, "description", "") or ""),
                "unit": str(getattr(item, "unit", "") or ""),
                "quantity": qty,
                "material": material,
                "skilled": skilled,
                "unskilled": unskilled,
                "labour": labour,
                "total": total,
                "unit_rate": unit_rate,
                "analysis": analysis,
            })
        return rows

    def export_excel(self, project, items, analyses, file_path=None):
        if project is None:
            raise ValueError("No active project selected.")

        code = str(getattr(project, "project_code", "") or "PROJECT")
        path = Path(file_path) if file_path else (
            self.export_folder / f"{code}_Rate_Analysis.xlsx"
        )
        path.parent.mkdir(parents=True, exist_ok=True)

        rows = self._rows(items, analyses)
        wb = Workbook()
        ws = wb.active
        ws.title = "Rate Analysis"
        ws.sheet_view.showGridLines = False

        ws.merge_cells("A1:K1")
        ws["A1"] = "RATE ANALYSIS / UNIT RATE BUILD-UP"
        ws["A1"].font = Font(size=17, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")

        ws["A3"] = "Project"
        ws["B3"] = str(getattr(project, "project_name", "") or "")
        ws["A4"] = "Project Code"
        ws["B4"] = code
        ws["A5"] = "Client"
        ws["B5"] = str(getattr(project, "client_name", "") or "")
        for r in range(3, 6):
            ws.cell(r, 1).font = Font(bold=True)

        headers = [
            "S.No.", "Item No.", "Description", "Unit", "Quantity",
            "Material Cost", "Skilled Labour", "Unskilled Labour",
            "Total Labour", "Total Cost", "Unit Rate"
        ]
        row = 7
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

        thin = Side(style="thin", color="C8C8C8")
        row += 1
        totals = [0.0] * 6
        for i, data in enumerate(rows, 1):
            vals = [
                i, data["item_no"], data["description"], data["unit"],
                data["quantity"], data["material"], data["skilled"],
                data["unskilled"], data["labour"], data["total"],
                data["unit_rate"],
            ]
            for c, value in enumerate(vals, 1):
                cell = ws.cell(row, c, value)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(
                    horizontal="right" if c >= 5 else (
                        "center" if c in (1, 4) else "left"
                    ),
                    vertical="top",
                    wrap_text=True,
                )
            for c in range(5, 12):
                ws.cell(row, c).number_format = '#,##0.00'
            totals[0] += data["quantity"]
            totals[1] += data["material"]
            totals[2] += data["skilled"]
            totals[3] += data["unskilled"]
            totals[4] += data["labour"]
            totals[5] += data["total"]
            row += 1

        ws.cell(row, 1, "GRAND TOTAL").font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        ws.cell(row, 5, totals[0])
        ws.cell(row, 6, totals[1])
        ws.cell(row, 7, totals[2])
        ws.cell(row, 8, totals[3])
        ws.cell(row, 9, totals[4])
        ws.cell(row, 10, totals[5])
        for c in range(5, 11):
            ws.cell(row, c).font = Font(bold=True)
            ws.cell(row, c).number_format = '#,##0.00'

        widths = [8, 15, 34, 10, 15, 18, 18, 18, 18, 18, 18]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A8"
        ws.auto_filter.ref = f"A7:K{max(row - 1, 7)}"

        # Detailed build-up sheet
        detail = wb.create_sheet("Cost Build-up")
        detail.sheet_view.showGridLines = False
        detail.merge_cells("A1:G1")
        detail["A1"] = "DETAILED COST BUILD-UP"
        detail["A1"].font = Font(size=16, bold=True)
        detail["A1"].alignment = Alignment(horizontal="center")

        r = 3
        for data in rows:
            detail.cell(r, 1, data["item_no"]).font = Font(size=12, bold=True)
            detail.cell(r, 2, data["description"]).font = Font(size=12, bold=True)
            r += 1
            detail_headers = ["Component", "Name", "Unit", "Quantity", "Rate", "Amount", "Type"]
            for c, h in enumerate(detail_headers, 1):
                cell = detail.cell(r, c, h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="5B7FA3")
                cell.alignment = Alignment(horizontal="center")
            r += 1

            for mat in data["analysis"].get("materials", []) or []:
                name = str(mat.get("name") or mat.get("material_name") or "Material")
                qty = self._num(mat.get("quantity", mat.get("QTY")))
                rate = self._num(mat.get("rate", mat.get("RATE")))
                amount = self._num(mat.get("amount", mat.get("cost", qty * rate)))
                vals = ["Material", name, str(mat.get("unit", mat.get("UNIT", ""))), qty, rate, amount, "Material"]
                for c, v in enumerate(vals, 1): detail.cell(r, c, v)
                r += 1

            for lab in data["analysis"].get("labour", []) or []:
                name = str(lab.get("name") or lab.get("labour") or "Labour")
                qty = self._num(lab.get("quantity", lab.get("days")))
                rate = self._num(lab.get("rate"))
                amount = self._num(lab.get("amount", qty * rate))
                vals = ["Labour", name, str(lab.get("unit", "day")), qty, rate, amount, "Labour"]
                for c, v in enumerate(vals, 1): detail.cell(r, c, v)
                r += 1

            r += 1

        for col, width in enumerate([15, 32, 12, 15, 18, 20, 15], 1):
            detail.column_dimensions[get_column_letter(col)].width = width

        wb.save(path)
        return path

    def export_pdf(self, project, items, analyses, file_path=None):
        if project is None:
            raise ValueError("No active project selected.")

        code = str(getattr(project, "project_code", "") or "PROJECT")
        path = Path(file_path) if file_path else (
            self.export_folder / f"{code}_Rate_Analysis.pdf"
        )
        path.parent.mkdir(parents=True, exist_ok=True)

        rows = self._rows(items, analyses)
        doc = SimpleDocTemplate(
            str(path), pagesize=landscape(A4),
            rightMargin=10*mm, leftMargin=10*mm,
            topMargin=10*mm, bottomMargin=10*mm,
        )
        styles = getSampleStyleSheet()
        title = ParagraphStyle(
            "RateTitle", parent=styles["Title"],
            alignment=TA_CENTER, fontSize=17, leading=20
        )
        small = ParagraphStyle(
            "RateSmall", parent=styles["BodyText"],
            fontSize=7.5, leading=9
        )

        story = [
            Paragraph("RATE ANALYSIS / UNIT RATE BUILD-UP", title),
            Spacer(1, 3*mm),
            Paragraph(
                f"<b>Project:</b> {getattr(project,'project_name','')} &nbsp;&nbsp; "
                f"<b>Code:</b> {code} &nbsp;&nbsp; "
                f"<b>Client:</b> {getattr(project,'client_name','')}",
                small
            ),
            Spacer(1, 4*mm),
        ]

        data = [[
            "S.No.", "Item", "Description", "Unit", "Qty",
            "Material", "Skilled", "Unskilled", "Labour", "Total", "Unit Rate"
        ]]
        for i, x in enumerate(rows, 1):
            data.append([
                str(i), x["item_no"], x["description"], x["unit"],
                f"{x['quantity']:,.3f}", f"{x['material']:,.2f}",
                f"{x['skilled']:,.2f}", f"{x['unskilled']:,.2f}",
                f"{x['labour']:,.2f}", f"{x['total']:,.2f}",
                f"{x['unit_rate']:,.2f}",
            ])
        if len(data) == 1:
            data.append(["", "", "No analysed BOQ items", "", "", "", "", "", "", "", ""])

        table = Table(
            data,
            colWidths=[11*mm, 22*mm, 48*mm, 13*mm, 18*mm,
                       23*mm, 23*mm, 23*mm, 23*mm, 25*mm, 24*mm],
            repeatRows=1,
        )
        table.setStyle(TableStyle([
            ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E78")),
            ("TEXTCOLOR",(0,0),(-1,0),colors.white),
            ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
            ("FONTSIZE",(0,0),(-1,-1),7),
            ("ALIGN",(0,0),(0,-1),"CENTER"),
            ("ALIGN",(3,1),(-1,-1),"RIGHT"),
            ("GRID",(0,0),(-1,-1),0.35,colors.grey),
            ("VALIGN",(0,0),(-1,-1),"TOP"),
        ]))
        story.append(table)
        story.append(Spacer(1, 4*mm))
        story.append(Paragraph(
            "Unit Rate = Total Cost ÷ BOQ Quantity. "
            "Material and labour components are taken from the saved calculator analysis.",
            small
        ))
        story.append(Paragraph(
            f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}",
            small
        ))
        doc.build(story)
        return path


__all__ = ["RateAnalysisReportService"]
