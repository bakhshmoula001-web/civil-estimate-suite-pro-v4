"""
Civil Estimate Suite Pro v4.0
Detailed Estimate Report Service

Generates a complete engineer-friendly Excel/PDF report from the
saved BOQ calculator analyses.

Report flow:
BOQ item -> Dimensions/Quantity -> Materials -> Labour -> Costs -> Totals
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from services.labour_analysis_service import LabourAnalysisService

from reportlab.platypus import (
    Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    PageBreak,
)


class MaterialReportService:
    """Generate consolidated material, labour and detailed cost reports."""

    def __init__(self, export_folder: str | Path = "exports"):
        self.export_folder = Path(export_folder)
        self.export_folder.mkdir(parents=True, exist_ok=True)

    @staticmethod
    def _num(value, default=0.0):
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return default

    @staticmethod
    def _safe_project_value(project, *names) -> str:
        for name in names:
            value = getattr(project, name, None)
            if value not in (None, ""):
                return str(value)
        return ""

    @classmethod
    def aggregate(cls, analyses) -> list[dict]:
        """Aggregate material quantities for the dashboard.

        Dashboard analysis data is already stored per BOQ item, so this
        method intentionally works directly on that analysis mapping/list.
        It returns the compact row format expected by DashboardPage:
        name, quantity and unit.
        """
        materials = defaultdict(float)

        if isinstance(analyses, dict):
            analysis_values = analyses.values()
        elif isinstance(analyses, (list, tuple)):
            analysis_values = analyses
        else:
            analysis_values = []

        for analysis in analysis_values:
            if not isinstance(analysis, dict):
                continue

            for entry in analysis.get("materials", []) or []:
                if not isinstance(entry, dict):
                    continue

                name = str(
                    entry.get("name", "Material")
                ).strip() or "Material"
                unit = str(
                    entry.get("unit", "")
                ).strip()

                materials[(name, unit)] += cls._num(
                    entry.get("quantity")
                )

        return [
            {
                "name": name,
                "quantity": round(quantity, 3),
                "unit": unit,
            }
            for (name, unit), quantity
            in sorted(materials.items())
        ]

    @classmethod
    def consolidate(cls, items: Iterable, analyses: dict[int, dict] | None = None):
        analyses = analyses or {}
        materials = defaultdict(float)
        labour = defaultdict(float)

        for item in items:
            analysis = analyses.get(getattr(item, "id", None), {})
            if not isinstance(analysis, dict):
                continue

            for entry in analysis.get("materials", []) or []:
                name = str(entry.get("name", "Material")).strip() or "Material"
                unit = str(entry.get("unit", "")).strip()
                materials[(name, unit)] += cls._num(entry.get("quantity"))

            labour_analysis = LabourAnalysisService.analyze(
                analysis.get("labour", []) or []
            )
            for entry in labour_analysis["rows"]:
                name = entry["name"]
                unit = entry["unit"]
                labour[(name, unit)] += entry["quantity"]

        material_rows = [
            {"name": name, "quantity": round(qty, 3), "unit": unit}
            for (name, unit), qty in sorted(materials.items())
        ]
        labour_rows = [
            {"name": name, "quantity": round(qty, 2), "unit": unit}
            for (name, unit), qty in sorted(labour.items())
        ]
        return material_rows, labour_rows

    @classmethod
    def _detail_rows(cls, items, analyses):
        rows = []
        for serial, item in enumerate(items, 1):
            analysis = analyses.get(getattr(item, "id", None), {})
            if not isinstance(analysis, dict):
                analysis = {}

            dimensions = analysis.get("dimensions", {}) or {}
            dim_unit = dimensions.get("dimension_unit", "")
            dim_text = ""
            if isinstance(dimensions, dict):
                parts = []
                for key, label in (("length", "L"), ("width", "W"), ("height", "T")):
                    if dimensions.get(key) is not None:
                        parts.append(f"{label}={cls._num(dimensions.get(key)):,.3f}")
                if parts:
                    dim_text = " × ".join(parts) + (f" {dim_unit}" if dim_unit else "")

            material_total = cls._num(analysis.get("material_total"))
            labour_analysis = LabourAnalysisService.analyze(
                analysis.get("labour", []) or []
            )
            labour_total = labour_analysis["labour_total"]
            total_cost = cls._num(analysis.get("total_cost"))
            if total_cost == 0:
                total_cost = material_total + labour_total

            rows.append({
                "sno": serial,
                "item_no": getattr(item, "item_no", ""),
                "description": getattr(item, "description", ""),
                "dimensions": dim_text,
                "unit": getattr(item, "unit", ""),
                "quantity": cls._num(getattr(item, "quantity", 0)),
                "materials": analysis.get("materials", []) or [],
                "labour": labour_analysis["rows"],
                "material_total": material_total,
                "labour_total": labour_total,
                "total_cost": total_cost,
                "unit_rate": cls._num(analysis.get("unit_rate")),
                "calculator": analysis.get("_calculator_type", ""),
            })
        return rows

    @staticmethod
    def _style_table(ws, start_row, end_row, start_col, end_col, header_row=None):
        thin = Side(style="thin", color="B7B7B7")
        for row in ws.iter_rows(
            min_row=start_row, max_row=end_row,
            min_col=start_col, max_col=end_col
        ):
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        if header_row:
            fill = PatternFill("solid", fgColor="1F4E78")
            font = Font(bold=True, color="FFFFFF")
            for cell in ws[header_row][start_col-1:end_col]:
                cell.fill = fill
                cell.font = font
                cell.alignment = Alignment(horizontal="center", vertical="center")

    def export_excel(self, project, items, analyses=None) -> str:
        analyses = analyses or {}
        details = self._detail_rows(items, analyses)
        material_rows, labour_rows = self.consolidate(items, analyses)

        wb = Workbook()
        ws = wb.active
        ws.title = "Detailed Estimate"
        ws.freeze_panes = "A8"

        ws.merge_cells("A1:L1")
        ws["A1"] = "DETAILED ESTIMATE REPORT"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A1"].alignment = Alignment(horizontal="center")

        info = [
            ("Project Code", self._safe_project_value(project, "project_code", "code")),
            ("Project Name", self._safe_project_value(project, "project_name", "name")),
            ("Client Name", self._safe_project_value(project, "client_name", "client")),
            ("Location", self._safe_project_value(project, "location")),
        ]
        for r, (label, value) in enumerate(info, 3):
            ws.cell(r, 1, label).font = Font(bold=True)
            ws.cell(r, 2, value)
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)

        headers = [
            "S.No", "Item No", "Work Item", "Dimensions", "Unit", "Quantity",
            "Material Quantity", "Skilled Labour", "Unskilled Labour",
            "Material Cost (PKR)", "Labour Cost (PKR)", "Total Cost (PKR)"
        ]
        row = 8
        for col, value in enumerate(headers, 1):
            ws.cell(row, col, value)
        self._style_table(ws, row, row, 1, 12, header_row=row)

        row += 1
        grand_material = grand_labour = grand_total = 0.0
        for d in details:
            material_text = "\n".join(
                f"{x.get('name','Material')}: {self._num(x.get('quantity')):,.3f} {x.get('unit','')}"
                for x in d["materials"]
            ) or "—"
            skilled = []
            unskilled = []
            for x in d["labour"]:
                text = f"{self._num(x.get('quantity')):,.2f} {x.get('unit','day')} @ PKR {self._num(x.get('rate')):,.2f}"
                if "unskilled" in str(x.get("name","")).lower():
                    unskilled.append(text)
                else:
                    skilled.append(text)

            values = [
                d["sno"], d["item_no"], d["description"], d["dimensions"],
                d["unit"], d["quantity"], material_text,
                "\n".join(skilled) or "—", "\n".join(unskilled) or "—",
                d["material_total"], d["labour_total"], d["total_cost"]
            ]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
            row += 1
            grand_material += d["material_total"]
            grand_labour += d["labour_total"]
            grand_total += d["total_cost"]

        ws.cell(row, 9, "TOTAL")
        ws.cell(row, 10, grand_material)
        ws.cell(row, 11, grand_labour)
        ws.cell(row, 12, grand_total)
        for col in range(9, 13):
            ws.cell(row, col).font = Font(bold=True)
        self._style_table(ws, 8, row, 1, 12)

        for r in range(9, row + 1):
            ws.cell(r, 6).number_format = "#,##0.000"
            for c in (10, 11, 12):
                ws.cell(r, c).number_format = '#,##0.00'

        widths = [7, 14, 24, 30, 10, 14, 34, 26, 26, 20, 20, 20]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width

        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4)

        # Material summary
        ms = wb.create_sheet("Material Summary")
        ms.append(["S.No", "Material", "Quantity", "Unit"])
        for i, x in enumerate(material_rows, 1):
            ms.append([i, x["name"], x["quantity"], x["unit"]])
        self._style_table(ms, 1, max(1, ms.max_row), 1, 4, header_row=1)
        ms.freeze_panes = "A2"
        for c,w in enumerate([8,35,18,15],1):
            ms.column_dimensions[get_column_letter(c)].width=w

        # Labour summary
        ls = wb.create_sheet("Labour Summary")
        ls.append(["S.No", "Labour", "Quantity", "Unit"])
        for i, x in enumerate(labour_rows, 1):
            ls.append([i, x["name"], x["quantity"], x["unit"]])
        self._style_table(ls, 1, max(1, ls.max_row), 1, 4, header_row=1)
        ls.freeze_panes = "A2"
        for c,w in enumerate([8,35,18,15],1):
            ls.column_dimensions[get_column_letter(c)].width=w

        # Cost summary
        cs = wb.create_sheet("Cost Summary")
        cs.append(["Cost Head", "Amount (PKR)"])
        cs.append(["Material Cost", grand_material])
        cs.append(["Labour Cost", grand_labour])
        cs.append(["Grand Total", grand_total])
        self._style_table(cs, 1, 4, 1, 2, header_row=1)
        for r in range(2,5):
            cs.cell(r,2).number_format="#,##0.00"
        cs.column_dimensions["A"].width=28
        cs.column_dimensions["B"].width=22

        path = self.export_folder / "Detailed_Estimate_Report.xlsx"
        wb.save(path)
        return str(path)

    def export_pdf(self, project, items, analyses=None) -> str:
        analyses = analyses or {}
        details = self._detail_rows(items, analyses)
        path = self.export_folder / "Detailed_Estimate_Report.pdf"

        styles = getSampleStyleSheet()
        title = ParagraphStyle("DETTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                               fontSize=16, leading=19, alignment=TA_CENTER, spaceAfter=7)
        normal = ParagraphStyle("DETNormal", parent=styles["Normal"], fontSize=7.5, leading=9)
        header = ParagraphStyle("DETHeader", parent=normal, fontName="Helvetica-Bold",
                                textColor=colors.white, alignment=TA_CENTER)
        right = ParagraphStyle("DETRight", parent=normal, alignment=TA_RIGHT)

        doc = SimpleDocTemplate(str(path), pagesize=landscape(A4),
                                leftMargin=10*mm, rightMargin=10*mm,
                                topMargin=9*mm, bottomMargin=12*mm)

        story=[Paragraph("DETAILED ESTIMATE REPORT",title)]
        project_data=[
            [Paragraph("<b>Project Code</b>",normal),Paragraph(self._safe_project_value(project,"project_code","code"),normal),
             Paragraph("<b>Project Name</b>",normal),Paragraph(self._safe_project_value(project,"project_name","name"),normal)],
            [Paragraph("<b>Client Name</b>",normal),Paragraph(self._safe_project_value(project,"client_name","client"),normal),
             Paragraph("<b>Location</b>",normal),Paragraph(self._safe_project_value(project,"location"),normal)]
        ]
        info=Table(project_data,colWidths=[27*mm,68*mm,27*mm,145*mm])
        info.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.35,colors.HexColor("#B7B7B7")),
                                  ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                                  ("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),
                                  ("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
        story += [info, Spacer(1,6)]

        for d in details:
            story.append(Paragraph(
                f"<b>{d['sno']}. {d['item_no']} — {d['description']}</b> | "
                f"Dimensions: {d['dimensions'] or '—'} | Quantity: {d['quantity']:,.3f} {d['unit']}",
                normal))
            data=[[Paragraph("Type",header),Paragraph("Description",header),
                   Paragraph("Quantity",header),Paragraph("Unit",header),
                   Paragraph("Rate (PKR)",header),Paragraph("Amount (PKR)",header)]]
            for x in d["materials"]:
                data.append([Paragraph("Material",normal),Paragraph(str(x.get("name","")),normal),
                             Paragraph(f"{self._num(x.get('quantity')):,.3f}",right),
                             Paragraph(str(x.get("unit","")),normal),
                             Paragraph(f"{self._num(x.get('rate')):,.2f}",right),
                             Paragraph(f"{self._num(x.get('amount')):,.2f}",right)])
            for x in d["labour"]:
                data.append([Paragraph("Labour",normal),Paragraph(str(x.get("name","")),normal),
                             Paragraph(f"{self._num(x.get('quantity')):,.2f}",right),
                             Paragraph(str(x.get("unit","day")),normal),
                             Paragraph(f"{self._num(x.get('rate')):,.2f}",right),
                             Paragraph(f"{self._num(x.get('amount')):,.2f}",right)])
            data += [
                [Paragraph("",normal),Paragraph("<b>Material Total</b>",normal),"","","",
                 Paragraph(f"<b>{d['material_total']:,.2f}</b>",right)],
                [Paragraph("",normal),Paragraph("<b>Labour Total</b>",normal),"","","",
                 Paragraph(f"<b>{d['labour_total']:,.2f}</b>",right)],
                [Paragraph("",normal),Paragraph("<b>Total Cost</b>",normal),"","","",
                 Paragraph(f"<b>{d['total_cost']:,.2f}</b>",right)],
            ]
            table=Table(data,colWidths=[22*mm,75*mm,31*mm,25*mm,35*mm,40*mm],repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1F4E78")),
                ("GRID",(0,0),(-1,-1),.3,colors.HexColor("#B7B7B7")),
                ("VALIGN",(0,0),(-1,-1),"TOP"),
                ("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3),
                ("TOPPADDING",(0,0),(-1,-1),3),("BOTTOMPADDING",(0,0),(-1,-1),3),
                ("BACKGROUND",(0,-3),(-1,-1),colors.HexColor("#F2F2F2")),
            ]))
            story += [table, Spacer(1,6)]

        total_material=sum(d["material_total"] for d in details)
        total_labour=sum(d["labour_total"] for d in details)
        total_cost=sum(d["total_cost"] for d in details)
        summary=Table([
            [Paragraph("<b>PROJECT MATERIAL COST</b>",normal),Paragraph(f"<b>PKR {total_material:,.2f}</b>",right)],
            [Paragraph("<b>PROJECT LABOUR COST</b>",normal),Paragraph(f"<b>PKR {total_labour:,.2f}</b>",right)],
            [Paragraph("<b>PROJECT GRAND TOTAL</b>",normal),Paragraph(f"<b>PKR {total_cost:,.2f}</b>",right)],
        ],colWidths=[80*mm,45*mm])
        summary.setStyle(TableStyle([("GRID",(0,0),(-1,-1),.4,colors.HexColor("#777777")),
                                     ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#F2F2F2")),
                                     ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5),
                                     ("TOPPADDING",(0,0),(-1,-1),5),("BOTTOMPADDING",(0,0),(-1,-1),5)]))
        story += [Spacer(1,5),summary]

        story.append(Spacer(1,7))
        story.append(Paragraph(
            f"Generated by Civil Estimate Suite Pro | {datetime.now().strftime('%d-%m-%Y %H:%M')}",
            ParagraphStyle("Footer",parent=normal,fontSize=7,alignment=TA_RIGHT)))

        doc.build(story)
        return str(path)

    def export_both(self, project, items, analyses=None):
        return self.export_excel(project, items, analyses), self.export_pdf(project, items, analyses)


__all__=["MaterialReportService"]
