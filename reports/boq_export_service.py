"""
=========================================================
Civil Estimate Suite Pro v4.0
---------------------------------------------------------
Module    : BOQ Export Service
Purpose   : Professional Excel / PDF BOQ Reports
Version   : 4.0.0
=========================================================
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class BOQExportService:
    """
    Generates professional BOQ reports from project + BOQ model objects.

    The service is deliberately independent from GUI code.
    """

    def __init__(self, export_folder: str | Path = "exports"):
        self.export_folder = Path(export_folder)

    # =====================================================
    # PUBLIC API
    # =====================================================

    def export_excel(
        self,
        project,
        items: Iterable,
        file_path: str | Path | None = None,
    ) -> Path:
        """Create a formatted Excel BOQ report."""
        project = self._require_project(project)
        items = list(items)

        path = self._resolve_path(
            file_path,
            default_name="BOQ_Report.xlsx",
        )

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "BOQ"

        self._build_excel(
            sheet,
            project,
            items,
        )

        workbook.save(path)
        return path

    def export_pdf(
        self,
        project,
        items: Iterable,
        file_path: str | Path | None = None,
    ) -> Path:
        """Create a professional landscape A4 PDF BOQ report."""
        project = self._require_project(project)
        items = list(items)

        path = self._resolve_path(
            file_path,
            default_name="BOQ_Report.pdf",
        )

        document = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            rightMargin=12 * mm,
            leftMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title="Bill of Quantities",
            author="Civil Estimate Suite Pro",
        )

        story = self._build_pdf_story(
            project,
            items,
        )

        document.build(
            story,
            onFirstPage=self._pdf_page,
            onLaterPages=self._pdf_page,
        )

        return path

    def export_both(
        self,
        project,
        items: Iterable,
        folder: str | Path | None = None,
    ) -> tuple[Path, Path]:
        """Generate both Excel and PDF reports."""
        target = (
            Path(folder)
            if folder is not None
            else self.export_folder
        )

        target.mkdir(
            parents=True,
            exist_ok=True,
        )

        excel_path = target / "BOQ_Report.xlsx"
        pdf_path = target / "BOQ_Report.pdf"

        return (
            self.export_excel(
                project,
                items,
                excel_path,
            ),
            self.export_pdf(
                project,
                items,
                pdf_path,
            ),
        )

    # =====================================================
    # COMMON HELPERS
    # =====================================================

    def _require_project(self, project):
        if project is None:
            raise ValueError(
                "A project is required for BOQ export."
            )

        if getattr(project, "id", None) is None:
            raise ValueError(
                "The current project has no valid ID."
            )

        return project

    def _resolve_path(
        self,
        file_path,
        default_name: str,
    ) -> Path:
        if file_path is None:
            path = (
                self.export_folder
                / default_name
            )
        else:
            path = Path(file_path)

        path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        return path

    @staticmethod
    def _project_value(
        project,
        field: str,
        default: str = "",
    ) -> str:
        value = getattr(
            project,
            field,
            default,
        )

        if value is None:
            return default

        return str(value).strip()

    @staticmethod
    def _item_amount(item) -> float:
        try:
            return round(
                float(
                    getattr(
                        item,
                        "amount",
                        0.0,
                    )
                    or 0.0
                ),
                2,
            )
        except (
            TypeError,
            ValueError,
        ):
            quantity = float(
                getattr(
                    item,
                    "quantity",
                    0.0,
                )
                or 0.0
            )

            rate = float(
                getattr(
                    item,
                    "rate",
                    0.0,
                )
                or 0.0
            )

            return round(
                quantity * rate,
                2,
            )

    # =====================================================
    # EXCEL
    # =====================================================

    def _build_excel(
        self,
        sheet,
        project,
        items,
    ) -> None:
        sheet.sheet_view.showGridLines = False

        title = "BILL OF QUANTITIES"

        sheet.merge_cells(
            "A1:H1"
        )

        sheet["A1"] = title
        sheet["A1"].font = Font(
            bold=True,
            size=18,
        )
        sheet["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        sheet.row_dimensions[1].height = 28

        project_fields = [
            (
                "Project Code",
                self._project_value(
                    project,
                    "project_code",
                ),
            ),
            (
                "Project Name",
                self._project_value(
                    project,
                    "project_name",
                ),
            ),
            (
                "Client Name",
                self._project_value(
                    project,
                    "client_name",
                ),
            ),
            (
                "Location",
                self._project_value(
                    project,
                    "location",
                ),
            ),
        ]

        row = 3

        for label, value in project_fields:
            sheet.cell(
                row=row,
                column=1,
                value=label,
            ).font = Font(
                bold=True,
            )

            sheet.cell(
                row=row,
                column=2,
                value=value,
            )

            sheet.merge_cells(
                start_row=row,
                start_column=2,
                end_row=row,
                end_column=8,
            )

            row += 1

        row += 1

        headers = [
            "S.No.",
            "Item No.",
            "Description",
            "Unit",
            "Quantity",
            "Rate (PKR)",
            "Amount (PKR)",
            "Remarks",
        ]

        header_row = row

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        header_font = Font(
            bold=True,
            color="FFFFFF",
        )

        thin = Side(
            style="thin",
            color="B7B7B7",
        )

        for column, value in enumerate(
            headers,
            start=1,
        ):
            cell = sheet.cell(
                row=header_row,
                column=column,
                value=value,
            )

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            cell.border = Border(
                top=thin,
                bottom=thin,
                left=thin,
                right=thin,
            )

        row += 1
        total = 0.0

        for serial, item in enumerate(
            items,
            start=1,
        ):
            quantity = float(
                getattr(
                    item,
                    "quantity",
                    0.0,
                )
                or 0.0
            )

            rate = float(
                getattr(
                    item,
                    "rate",
                    0.0,
                )
                or 0.0
            )

            amount = self._item_amount(
                item
            )

            total += amount

            values = [
                serial,
                getattr(
                    item,
                    "item_no",
                    "",
                ),
                getattr(
                    item,
                    "description",
                    "",
                ),
                getattr(
                    item,
                    "unit",
                    "",
                ),
                quantity,
                rate,
                amount,
                getattr(
                    item,
                    "remarks",
                    "",
                ),
            ]

            for column, value in enumerate(
                values,
                start=1,
            ):
                cell = sheet.cell(
                    row=row,
                    column=column,
                    value=value,
                )

                cell.border = Border(
                    top=thin,
                    bottom=thin,
                    left=thin,
                    right=thin,
                )

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal=(
                        "left"
                        if column
                        in (2, 3, 8)
                        else "center"
                    ),
                )

            sheet.cell(
                row=row,
                column=5,
            ).number_format = '#,##0.000'

            sheet.cell(
                row=row,
                column=6,
            ).number_format = '#,##0.00'

            sheet.cell(
                row=row,
                column=7,
            ).number_format = '#,##0.00'

            row += 1

        total_row = row

        sheet.merge_cells(
            start_row=total_row,
            start_column=1,
            end_row=total_row,
            end_column=6,
        )

        total_label = sheet.cell(
            row=total_row,
            column=1,
            value="GRAND TOTAL",
        )

        total_label.font = Font(
            bold=True,
            size=12,
        )

        total_label.alignment = Alignment(
            horizontal="right",
        )

        total_cell = sheet.cell(
            row=total_row,
            column=7,
            value=round(
                total,
                2,
            ),
        )

        total_cell.font = Font(
            bold=True,
            size=12,
        )

        total_cell.number_format = '#,##0.00'

        remarks_cell = sheet.cell(
            row=total_row,
            column=8,
            value="PKR",
        )

        remarks_cell.font = Font(
            bold=True,
        )

        for column in range(1, 9):
            sheet.cell(
                row=total_row,
                column=column,
            ).border = Border(
                top=thin,
                bottom=thin,
            )

        footer_row = total_row + 2

        sheet.merge_cells(
            start_row=footer_row,
            start_column=1,
            end_row=footer_row,
            end_column=8,
        )

        generated = datetime.now().strftime(
            "%d-%m-%Y %H:%M"
        )

        sheet.cell(
            row=footer_row,
            column=1,
            value=(
                "Generated by Civil Estimate Suite Pro "
                f"| {generated}"
            ),
        ).font = Font(
            italic=True,
            size=9,
        )

        sheet.cell(
            row=footer_row,
            column=1,
        ).alignment = Alignment(
            horizontal="right",
        )

        widths = [
            8,
            14,
            42,
            12,
            15,
            16,
            18,
            28,
        ]

        for index, width in enumerate(
            widths,
            start=1,
        ):
            sheet.column_dimensions[
                get_column_letter(index)
            ].width = width

        sheet.freeze_panes = (
            f"A{header_row + 1}"
        )

        sheet.auto_filter.ref = (
            f"A{header_row}:H{max(row - 1, header_row)}"
        )

        sheet.page_setup.orientation = (
            "landscape"
        )
        sheet.page_setup.paperSize = (
            sheet.PAPERSIZE_A4
        )
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0

        sheet.sheet_properties.pageSetUpPr.fitToPage = True

        sheet.oddFooter.center.text = (
            "Civil Estimate Suite Pro"
        )
        sheet.oddFooter.right.text = (
            "Page &P of &N"
        )

    # =====================================================
    # PDF
    # =====================================================

    def _build_pdf_story(
        self,
        project,
        items,
    ):
        """
        Build a professional landscape A4 BOQ PDF.

        The PDF uses the same source data and calculations as the
        Excel export.  The layout is intentionally wider than the
        earlier version so Description and Remarks remain readable.
        """
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "BOQTitle",
            parent=styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=17,
            leading=20,
            alignment=TA_CENTER,
            spaceAfter=7,
        )

        info_label_style = ParagraphStyle(
            "BOQInfoLabel",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=8.2,
            leading=10,
        )

        info_value_style = ParagraphStyle(
            "BOQInfoValue",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.2,
            leading=10,
        )

        cell_style = ParagraphStyle(
            "BOQCell",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=7.5,
            leading=9,
            spaceAfter=0,
            spaceBefore=0,
        )

        cell_header = ParagraphStyle(
            "BOQHeader",
            parent=cell_style,
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=9,
            textColor=colors.white,
            alignment=TA_CENTER,
        )

        cell_center = ParagraphStyle(
            "BOQCellCenter",
            parent=cell_style,
            alignment=TA_CENTER,
        )

        cell_right = ParagraphStyle(
            "BOQCellRight",
            parent=cell_style,
            alignment=TA_RIGHT,
        )

        total_label_style = ParagraphStyle(
            "BOQTotalLabel",
            parent=cell_style,
            fontName="Helvetica-Bold",
            alignment=TA_RIGHT,
        )

        total_value_style = ParagraphStyle(
            "BOQTotalValue",
            parent=cell_style,
            fontName="Helvetica-Bold",
            alignment=TA_RIGHT,
        )

        total_currency_style = ParagraphStyle(
            "BOQTotalCurrency",
            parent=cell_style,
            fontName="Helvetica-Bold",
            alignment=TA_CENTER,
        )

        story = [
            Paragraph(
                "BILL OF QUANTITIES",
                title_style,
            )
        ]

        # -------------------------------------------------
        # PROJECT INFORMATION
        # -------------------------------------------------

        project_code = self._project_value(
            project,
            "project_code",
        )
        project_name = self._project_value(
            project,
            "project_name",
        )
        client_name = self._project_value(
            project,
            "client_name",
        )
        location = self._project_value(
            project,
            "location",
        )

        info_data = [
            [
                Paragraph("Project Code", info_label_style),
                Paragraph(self._escape_text(project_code), info_value_style),
                Paragraph("Project Name", info_label_style),
                Paragraph(self._escape_text(project_name), info_value_style),
            ],
            [
                Paragraph("Client Name", info_label_style),
                Paragraph(self._escape_text(client_name), info_value_style),
                Paragraph("Location", info_label_style),
                Paragraph(self._escape_text(location), info_value_style),
            ],
        ]

        info_table = Table(
            info_data,
            colWidths=[
                28 * mm,
                72 * mm,
                28 * mm,
                145 * mm,
            ],
            hAlign="LEFT",
        )

        info_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7B7B7")),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F2F2")),
                    ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F2F2F2")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )

        story.append(info_table)
        story.append(Spacer(1, 8))

        # -------------------------------------------------
        # BOQ TABLE
        # -------------------------------------------------

        table_data = [
            [
                Paragraph("S.No.", cell_header),
                Paragraph("Item No.", cell_header),
                Paragraph("Description", cell_header),
                Paragraph("Unit", cell_header),
                Paragraph("Quantity", cell_header),
                Paragraph("Rate (PKR)", cell_header),
                Paragraph("Amount (PKR)", cell_header),
                Paragraph("Remarks", cell_header),
            ]
        ]

        total = 0.0

        for serial, item in enumerate(items, start=1):
            try:
                quantity = float(
                    getattr(item, "quantity", 0.0) or 0.0
                )
            except (TypeError, ValueError):
                quantity = 0.0

            try:
                rate = float(
                    getattr(item, "rate", 0.0) or 0.0
                )
            except (TypeError, ValueError):
                rate = 0.0

            amount = self._item_amount(item)
            total += amount

            table_data.append(
                [
                    Paragraph(str(serial), cell_center),
                    Paragraph(
                        self._escape_text(getattr(item, "item_no", "")),
                        cell_center,
                    ),
                    Paragraph(
                        self._escape_text(getattr(item, "description", "")),
                        cell_style,
                    ),
                    Paragraph(
                        self._escape_text(getattr(item, "unit", "")),
                        cell_center,
                    ),
                    Paragraph(f"{quantity:,.3f}", cell_right),
                    Paragraph(f"{rate:,.2f}", cell_right),
                    Paragraph(f"{amount:,.2f}", cell_right),
                    Paragraph(
                        self._escape_text(getattr(item, "remarks", "")),
                        cell_style,
                    ),
                ]
            )

        # Keep the total visually identical to the Excel report:
        # label in Rate column, value in Amount column, currency in Remarks.
        table_data.append(
            [
                "",
                "",
                "",
                "",
                "",
                Paragraph("GRAND TOTAL", total_label_style),
                Paragraph(f"{total:,.2f}", total_value_style),
                Paragraph("PKR", total_currency_style),
            ]
        )

        # Landscape A4 width = 297 mm. With 12 mm margins on both
        # sides the usable width is 273 mm. The columns intentionally
        # consume the full width to avoid the compressed look.
        boq_table = Table(
            table_data,
            repeatRows=1,
            colWidths=[
                14 * mm,  # S.No.
                22 * mm,  # Item No.
                70 * mm,  # Description
                16 * mm,  # Unit
                26 * mm,  # Quantity
                31 * mm,  # Rate
                37 * mm,  # Amount
                57 * mm,  # Remarks
            ],
            hAlign="LEFT",
        )

        boq_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7B7B7")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("ALIGN", (0, 0), (1, -1), "CENTER"),
                    ("ALIGN", (3, 1), (3, -1), "CENTER"),
                    ("ALIGN", (4, 1), (6, -1), "RIGHT"),
                    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F2F2F2")),
                    ("LINEABOVE", (0, -1), (-1, -1), 0.9, colors.black),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )

        story.append(boq_table)
        story.append(Spacer(1, 6))

        generated = datetime.now().strftime("%d-%m-%Y %H:%M")

        story.append(
            Paragraph(
                f"Generated by Civil Estimate Suite Pro | {generated}",
                ParagraphStyle(
                    "GeneratedText",
                    parent=styles["Normal"],
                    fontName="Helvetica-Oblique",
                    fontSize=7.2,
                    leading=9,
                    alignment=TA_RIGHT,
                    textColor=colors.HexColor("#333333"),
                ),
            )
        )

        return story

    @staticmethod
    def _escape_text(value) -> str:
        text = str(value if value is not None else "")

        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

    @staticmethod
    def _pdf_page(canvas, document) -> None:
        canvas.saveState()

        width, _ = landscape(A4)

        canvas.setStrokeColor(colors.HexColor("#D0D0D0"))
        canvas.setLineWidth(0.35)
        canvas.line(12 * mm, 10 * mm, width - 12 * mm, 10 * mm)

        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#555555"))

        canvas.drawString(
            12 * mm,
            6 * mm,
            "Civil Estimate Suite Pro",
        )

        canvas.drawRightString(
            width - 12 * mm,
            6 * mm,
            f"Page {document.page}",
        )

        canvas.restoreState()


__all__ = [
    "BOQExportService",
]
